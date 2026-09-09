"""The workbook: the operator's actual product.

Everything upstream -- the search budget, the geocoder, the niche taxonomy, the scorer --
exists to fill the rows of this file. The operator works from the sheet, not from Postgres,
so the sheet is where correctness finally has to show up.

Three properties this module is responsible for, none of them cosmetic:

*Provenance.* Every evidence cell carries its source URL as a comment. The number in that
cell is about to be said out loud to a business owner ("you've got 340 reviews averaging
4.6"), and being wrong in that sentence costs the meeting. One right-click is the whole
distance between a number and its receipt. A cell whose source is unknown gets no comment
rather than a plausible one -- see `view.GOOGLE_EVIDENCE_COLUMNS`.

*A safe write.* The file is built under a temporary name in the destination directory and
moved into place with `os.replace`, which is atomic on both POSIX and Windows. Writing the
target directly means a crash, a full disk, or a `ValueError` on row 4,000 leaves a
truncated workbook where a good one used to be -- and the operator's marked-up copy is the
only record of a week of visits. Nothing here ever opens the destination for writing.

*Banding is read, not recomputed.* `audience_band` and `banding_method` come from the
`lead_bands` view (migration 0009), which ranks each business against its (niche, city)
cohort at read time. Recomputing a band here would produce a second answer that drifts from
the one every other reader sees, and the drift would be invisible: both would look
plausible.

The database is injected, not opened. `fetch_records` takes any DB-API connection and uses
`cursor.description` to name its columns, so this module imports no driver at all -- and a
test can hand it a fake cursor to check the row mapping without Postgres. The SQL itself is
exercised by the integration test in `tests/test_export.py`.
"""

from __future__ import annotations

import os
import tempfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Evidence
from .view import (
    COLUMNS,
    OPERATOR_COLUMNS,
    ContactPerson,
    ExportRow,
    LeadRecord,
    OperatorVerdict,
    build_rows,
    evidence_sources,
    niche_label,
)

# --- the query -----------------------------------------------------------------------

# Keys read out of provider payloads. `enrichments.data` is whatever the fetcher stored, so
# these names are a contract with a module that does not exist yet; they are listed here,
# once, so that wiring the enricher is a one-line change rather than a hunt through SQL,
# and `tests/test_export.py` pins them against a payload it writes itself.
GOOGLE_TYPE_KEY = "type"
GOOGLE_RATING_KEY = "rating"
GOOGLE_PEAK_HOURS_KEY = "peak_hours"
INSTAGRAM_FOLLOWERS_KEY = "followers"
INSTAGRAM_ENGAGEMENT_KEY = "engagement_rate"
INSTAGRAM_ADS_KEY = "runs_ads"

GOOGLE_SOURCE = "google_maps"
# Three separate real sources, not one -- `INSTAGRAM_SOURCE = "instagram"` was a placeholder
# for a module that did not exist yet when this query was written (see the module note
# above); none of the three modules that eventually shipped ever wrote that exact string,
# so the join it named matched nothing, silently, since the day each module landed.
# Proven live: a business with a real, manually-verified 107K-follower Instagram profile on
# file (`instagram_profile`) still exported with a blank `followers` cell, because the join
# was looking for `source = 'instagram'`.
INSTAGRAM_HANDLE_SOURCE = "instagram_handle"
INSTAGRAM_PROFILE_SOURCE = "instagram_profile"
ADS_SOURCE = "ad_library"

# One row per business, joined to the freshest of everything else.
#
# Notes on the shape:
#
# * `lead_bands` is the driving relation, so a business with no score, no enrichment and no
#   contact still appears -- the view LEFT JOINs those in, and the leads the pipeline failed
#   on are exactly the ones worth looking at.
# * `reviews` comes from the view rather than from the payload here, so the number printed
#   in the sheet is the same number the band was computed from. Parsing it a second time is
#   how a sheet ends up saying 'large' next to a blank review count.
# * `signals` needs a second look at `scores`, which `lead_bands` does not expose. Its
#   ORDER BY is a deliberate copy of the view's `latest_score` CTE: change one and the sheet
#   starts showing one score version's total beside another's signals. The integration test
#   inserts two versions and checks they agree.
# * JSON is extracted as text (`->>`) and parsed in Python. A SQL cast would be shorter and
#   would take the whole export down on one malformed payload; here a bad value blanks one
#   cell and the other 40,000 rows still print.
# * No ORDER BY. Row order is `view.sort_key`'s decision -- area, then score, with rules for
#   missing values -- and writing those rules a second time in SQL is how the two drift.
# * No WHERE. The sheet is the whole corpus; the autofilter on the header row is how the
#   operator narrows it to a niche or a neighbourhood, and it costs nothing to leave every
#   row available for that.
EXPORT_QUERY = f"""
SELECT lb.name                       AS business_name,
       lb.niche_id                   AS niche_id,
       g.data->>'{GOOGLE_TYPE_KEY}'  AS business_type,
       b.country                     AS country,
       b.state                       AS state,
       lb.city                       AS city,
       lb.search_area                AS search_area,
       b.address                     AS address,
       b.lat                         AS lat,
       b.lng                         AS lng,
       b.phone                       AS phone,
       b.email                       AS email,
       b.website                     AS website,
       -- `businesses.instagram_handle` is never written by anything today -- an enrichment
       -- finding is an observation, and nothing promotes one onto the business row (see
       -- `enrichment/service.py`'s own docstring on why). The real, current source of a
       -- found handle is the `instagram_handle` enrichment observation itself.
       coalesce(b.instagram_handle, ih.data->>'handle') AS instagram_handle,
       b.facebook_url                AS facebook_url,
       ct.name                       AS contact_name,
       ct.role                       AS contact_role,
       ct.phone                      AS contact_phone,
       ct.email                      AS contact_email,
       ct.source                     AS contact_source,
       lb.reviews                    AS reviews,
       g.data->>'{GOOGLE_RATING_KEY}'            AS rating,
       g.data->>'{GOOGLE_PEAK_HOURS_KEY}'        AS peak_hours,
       g.source_url                              AS google_source_url,
       ip.data->>'{INSTAGRAM_FOLLOWERS_KEY}'     AS followers,
       ip.data->>'{INSTAGRAM_ENGAGEMENT_KEY}'    AS engagement_rate,
       ad.data->>'{INSTAGRAM_ADS_KEY}'           AS runs_ads,
       coalesce(ip.source_url, ih.source_url)    AS instagram_source_url,
       lb.total                      AS total_score,
       lb.audience_band              AS audience_band,
       lb.banding_method             AS banding_method,
       sc.signals                    AS signals,
       -- Written by `Engine.execute_enrichment` into the latest `scores` row's own
       -- `evidence` jsonb, beside `pitch_angle` -- an internal, third-person briefing about
       -- the lead, not a message to send it, so it belongs with the agent's own notes on
       -- the lead rather than in `outreach` (whose `channel` column means "sent via", which
       -- an internal summary has none of). NULL on a `google-only` score, since enrichment
       -- has not run yet and there is nothing to summarise beyond the raw listing.
       sc.evidence->>'ai_summary'    AS ai_summary,
       wp.body                       AS website_pitch,
       autos.opportunity_ids         AS automation_opportunities,
       ap.body                       AS automation_pitch,
       v.my_verdict                  AS my_verdict,
       v.notes                       AS notes,
       v.contacted_on                AS contacted_on,
       v.channel                     AS channel,
       v.outcome                     AS outcome
  FROM lead_bands lb
  JOIN businesses b ON b.id = lb.business_id
  LEFT JOIN verdicts v ON v.business_id = b.id
  LEFT JOIN LATERAL (
      -- Most confident contact, most recent as the tiebreak. The operator asks for one
      -- name at the counter, so the sheet carries one.
      SELECT c.name, c.role, c.phone, c.email, c.source
        FROM contacts c
       WHERE c.business_id = b.id
       ORDER BY c.confidence DESC, c.found_at DESC, c.id DESC
       LIMIT 1
  ) ct ON TRUE
  LEFT JOIN LATERAL (
      SELECT e.data, e.source_url
        FROM enrichments e
       WHERE e.business_id = b.id AND e.source = '{GOOGLE_SOURCE}' AND e.status = 'ok'
       ORDER BY e.fetched_at DESC, e.id DESC
       LIMIT 1
  ) g ON TRUE
  LEFT JOIN LATERAL (
      SELECT e.data, e.source_url
        FROM enrichments e
       WHERE e.business_id = b.id AND e.source = '{INSTAGRAM_HANDLE_SOURCE}' AND e.status = 'ok'
       ORDER BY e.fetched_at DESC, e.id DESC
       LIMIT 1
  ) ih ON TRUE
  LEFT JOIN LATERAL (
      SELECT e.data, e.source_url
        FROM enrichments e
       WHERE e.business_id = b.id AND e.source = '{INSTAGRAM_PROFILE_SOURCE}' AND e.status = 'ok'
       ORDER BY e.fetched_at DESC, e.id DESC
       LIMIT 1
  ) ip ON TRUE
  LEFT JOIN LATERAL (
      SELECT e.data
        FROM enrichments e
       WHERE e.business_id = b.id AND e.source = '{ADS_SOURCE}' AND e.status = 'ok'
       ORDER BY e.fetched_at DESC, e.id DESC
       LIMIT 1
  ) ad ON TRUE
  LEFT JOIN LATERAL (
      SELECT s.signals, s.evidence
        FROM scores s
       WHERE s.business_id = b.id
       ORDER BY s.scored_at DESC, s.id DESC
       LIMIT 1
  ) sc ON TRUE
  LEFT JOIN LATERAL (
      SELECT array_agg(a.opportunity_id ORDER BY a.confidence DESC, a.opportunity_id)
                 AS opportunity_ids
        FROM automation_opportunities a
       WHERE a.business_id = b.id
  ) autos ON TRUE
  LEFT JOIN LATERAL (
      SELECT o.body FROM outreach o
       WHERE o.business_id = b.id AND o.kind = 'website_pitch'
       ORDER BY o.created_at DESC, o.id DESC LIMIT 1
  ) wp ON TRUE
  LEFT JOIN LATERAL (
      SELECT o.body FROM outreach o
       WHERE o.business_id = b.id AND o.kind = 'automation_pitch'
       ORDER BY o.created_at DESC, o.id DESC LIMIT 1
  ) ap ON TRUE
"""


def _as_int(value: Any) -> int | None:
    """Provider text -> int, or nothing. Never a zero standing in for a parse failure."""
    if value is None or isinstance(value, bool):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        pass
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def _as_float(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


TRUE_TOKENS = frozenset({"true", "t", "yes", "y", "1"})
FALSE_TOKENS = frozenset({"false", "f", "no", "n", "0"})


def _as_bool(value: Any) -> bool | None:
    """Unrecognised text is unknown, not False.

    `runs_ads` reaches the operator as an argument -- "you're paying for ads and sending
    that traffic to an Instagram bio" -- so a False invented from a typo in a payload is a
    claim made without evidence.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    token = str(value).strip().lower()
    if token in TRUE_TOKENS:
        return True
    if token in FALSE_TOKENS:
        return False
    return None


def _as_list(value: Any) -> tuple[str, ...]:
    """`scores.signals` and the automation aggregate, as a tuple of display strings.

    `signals` is jsonb and has been written as both a list and an object by different
    callers, so both are handled. Nothing is dropped on the floor: an unexpected shape is
    stringified and shown, because a signal the operator cannot see is a signal that might
    as well not have been detected.
    """
    if value is None:
        return ()
    if isinstance(value, str):
        return (value,) if value.strip() else ()
    if isinstance(value, Mapping):
        return tuple(f"{key}: {value[key]}" for key in sorted(value, key=str))
    if isinstance(value, Sequence):
        return tuple(str(item).strip() for item in value if str(item).strip())
    return (str(value),)


def record_from_row(row: Mapping[str, Any]) -> LeadRecord:
    """One row of `EXPORT_QUERY` -> one `LeadRecord`.

    Split out from `fetch_records` so the mapping can be tested against a hand-written row
    -- including the malformed payloads a live provider eventually sends -- without a
    database anywhere near it.
    """
    google_url = row.get("google_source_url")
    instagram_url = row.get("instagram_source_url")
    evidence = Evidence(
        reviews=_as_int(row.get("reviews")),
        rating=_as_float(row.get("rating")),
        followers=_as_int(row.get("followers")),
        engagement_rate=_as_float(row.get("engagement_rate")),
        runs_ads=_as_bool(row.get("runs_ads")),
    )
    return LeadRecord(
        business_name=row.get("business_name") or "",
        niche=niche_label(row.get("niche_id")),
        business_type=row.get("business_type"),
        country=row.get("country"),
        state=row.get("state"),
        city=row.get("city"),
        search_area=row.get("search_area"),
        address=row.get("address"),
        lat=_as_float(row.get("lat")),
        lng=_as_float(row.get("lng")),
        phone=row.get("phone"),
        email=row.get("email"),
        website=row.get("website"),
        instagram_handle=row.get("instagram_handle"),
        facebook_url=row.get("facebook_url"),
        contact=ContactPerson(
            name=row.get("contact_name"),
            role=row.get("contact_role"),
            phone=row.get("contact_phone"),
            email=row.get("contact_email"),
            source=row.get("contact_source"),
        ),
        evidence=evidence,
        peak_hours=row.get("peak_hours"),
        # The URLs the fetcher actually used, not a guess from the handle: `enrichments`
        # records where each payload came from, and that is the page the number is on.
        evidence_sources=evidence_sources(google_url=google_url, instagram_url=instagram_url),
        total_score=_as_int(row.get("total_score")),
        audience_band=row.get("audience_band"),
        banding_method=row.get("banding_method"),
        signals=_as_list(row.get("signals")),
        ai_summary=row.get("ai_summary"),
        website_pitch=row.get("website_pitch"),
        automation_opportunities=_as_list(row.get("automation_opportunities")),
        automation_pitch=row.get("automation_pitch"),
        verdict=OperatorVerdict(
            my_verdict=row.get("my_verdict"),
            notes=row.get("notes"),
            contacted_on=row.get("contacted_on"),
            channel=row.get("channel"),
            outcome=row.get("outcome"),
        ),
    )


def fetch_records(connection: Any) -> list[LeadRecord]:
    """Every exportable lead, as records. The connection is the caller's to open and close.

    Columns are read off `cursor.description` rather than assumed positionally, so
    reordering the SELECT list cannot silently shift `phone` into `email`.
    """
    with connection.cursor() as cursor:
        cursor.execute(EXPORT_QUERY)
        names = [description[0] for description in cursor.description]
        return [
            record_from_row(dict(zip(names, values, strict=True)))
            for values in cursor.fetchall()
        ]


# --- the workbook --------------------------------------------------------------------

SHEET_TITLE = "Leads"

# Shown by Excel as the comment's author. Not cosmetic: it tells the operator the note was
# written by the export rather than by a colleague marking up the sheet.
COMMENT_AUTHOR = "lead-engine"

# 8-digit ARGB. openpyxl left-pads a 6-digit value with '00' on write, so a colour declared
# as 'FFF2CC' reads back as '00FFF2CC' and any comparison against the constant fails.
HEADER_FILL_COLOR = "FFDCE6F1"
OPERATOR_HEADER_FILL_COLOR = "FFF2C744"
OPERATOR_FILL_COLOR = "FFFFF6D8"

DEFAULT_WIDTH = 18
COLUMN_WIDTHS: dict[str, int] = {
    "business_name": 34,
    "niche": 20,
    "business_type": 20,
    "country": 12,
    "state": 16,
    "city": 16,
    "search_area": 20,
    "address": 44,
    "lat": 12,
    "lng": 12,
    "phone": 16,
    "email": 26,
    "website": 34,
    "instagram_handle": 20,
    "facebook_url": 30,
    "contact_email": 26,
    "reviews": 10,
    "rating": 8,
    "followers": 12,
    "engagement_rate": 14,
    "runs_ads": 10,
    "peak_hours": 22,
    "total_score": 12,
    "audience_band": 14,
    "banding_method": 15,
    "signals": 40,
    "ai_summary": 50,
    "website_pitch": 50,
    "automation_opportunities": 34,
    "automation_pitch": 50,
    "my_verdict": 14,
    "notes": 40,
    "contacted_on": 14,
}

# Long prose. Wrapped and top-aligned so a row stays one row instead of one very wide line.
WRAPPED_COLUMNS = frozenset(
    {
        "address",
        "signals",
        "ai_summary",
        "website_pitch",
        "automation_opportunities",
        "automation_pitch",
        "notes",
    }
)

NUMBER_FORMATS: dict[str, str] = {
    "lat": "0.000000",
    "lng": "0.000000",
    "reviews": "#,##0",
    "rating": "0.0",
    "followers": "#,##0",
    # The value stays the fraction it is everywhere else in the system (0.042); only its
    # display is a percentage. Storing 4.2 under a header that says `engagement_rate` would
    # make the sheet disagree with the scorer.
    "engagement_rate": "0.0%",
    "total_score": "0",
    "contacted_on": "yyyy-mm-dd",
}


def build_workbook(rows: Iterable[ExportRow]) -> Workbook:
    """Rows in, workbook out. No file touched.

    `rows` is consumed lazily and exactly once, so a generator that fails partway -- a
    cursor whose connection drops on row 4,000 -- raises here, before anything has been
    moved into place.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL_COLOR)
    operator_header_fill = PatternFill("solid", fgColor=OPERATOR_HEADER_FILL_COLOR)
    operator_fill = PatternFill("solid", fgColor=OPERATOR_FILL_COLOR)
    wrapped = Alignment(wrap_text=True, vertical="top")
    plain = Alignment(vertical="top")

    for index, column in enumerate(COLUMNS, start=1):
        # The header is the column key verbatim, not a prettified label. A later phase reads
        # the operator's edits back by column name, and "My verdict" would have to be
        # un-prettified to find it again.
        cell = sheet.cell(row=1, column=index, value=column)
        cell.font = header_font
        cell.fill = operator_header_fill if column in OPERATOR_COLUMNS else header_fill
        cell.alignment = Alignment(vertical="center")
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = COLUMN_WIDTHS.get(column, DEFAULT_WIDTH)

    written = 0
    for row in rows:
        if len(row.cells) != len(COLUMNS):
            # The row builder guarantees width structurally; this catches a hand-built row,
            # and catches it before the file is written rather than after the operator
            # notices every value in a column is one row's worth of data to the left.
            raise ValueError(
                f"row has {len(row.cells)} cells, expected {len(COLUMNS)}"
            )
        written += 1
        excel_row = written + 1
        for index, (column, source_cell) in enumerate(zip(COLUMNS, row.cells, strict=True), 1):
            cell = sheet.cell(row=excel_row, column=index, value=source_cell.value)
            cell.alignment = wrapped if column in WRAPPED_COLUMNS else plain
            number_format = NUMBER_FORMATS.get(column)
            if number_format:
                cell.number_format = number_format
            if column in OPERATOR_COLUMNS:
                cell.fill = operator_fill
            if source_cell.source:
                # Provenance, one right-click away. Written wherever a source exists, which
                # in practice is the evidence group -- see `view.evidence_sources`.
                cell.comment = Comment(source_cell.source, COMMENT_AUTHOR)

    sheet.freeze_panes = "B2"
    # Column A is frozen along with the header: the operator columns are at the far right of
    # 40-odd columns, and scrolling to them without the business name in view is how a
    # verdict lands on the wrong business.
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{written + 1}"
    return workbook


def write_workbook(rows: Iterable[ExportRow], path: str | os.PathLike[str]) -> Path:
    """Write the sheet to `path`, atomically. Returns the path written.

    The temporary file is created in the destination directory, not in the system temp
    directory: `os.replace` is only atomic within a filesystem, and a cross-device move
    degrades to copy-then-delete, which is exactly the half-written window this avoids.

    If anything raises -- a bad value, a full disk, a dropped cursor -- the temporary file
    is removed and the previous sheet is left byte-for-byte as it was. One failure is not
    handled and must not be: if the operator has the workbook open in Excel, Windows refuses
    the replace and this raises `PermissionError` with the old file intact. That is the
    correct outcome; the alternative is overwriting a file someone is editing.
    """
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(
        dir=target.parent, prefix=f".{target.name}.", suffix=".tmp"
    )
    os.close(handle)
    temporary = Path(temporary_name)
    try:
        workbook = build_workbook(rows)
        workbook.save(temporary)
        # Durability before visibility. Without the fsync the rename can be on disk while
        # the bytes it points at are still in the page cache, and a power loss then leaves a
        # zero-length file where the old sheet was. 'r+b' rather than 'rb': on Windows
        # os.fsync needs a handle opened for writing.
        with open(temporary, "r+b") as stream:
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, target)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return target


def export_records(records: Iterable[LeadRecord], path: str | os.PathLike[str]) -> Path:
    """Records -> sorted rows -> workbook on disk."""
    return write_workbook(build_rows(records), path)


def export(connection: Any, path: str | os.PathLike[str]) -> Path:
    """The whole job: read the banded corpus, lay it out, write the sheet."""
    return export_records(fetch_records(connection), path)
